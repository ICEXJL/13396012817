# -*- coding: utf-8 -*-
"""组合表兜底、低库存调零和预售加量的匹配测试。"""

import os

os.environ["QT_QPA_PLATFORM"] = "offscreen"

from openpyxl import Workbook

from match_stock_gui import AutoMatchWorker


def _sales_file(path):
    book = Workbook()
    sheet = book.active
    sheet.append(["商品ID", "商品名称", "规格编号", "主商品货号", "商品货号", "卖家库存"])
    sheet.append(["1", "普通低库存", "1", "", "10000001-red", 8])
    sheet.append(["2", "组合预售", "700002", "", "zuhe-special", 3])
    book.save(path)


def _stock_file(path):
    book = Workbook()
    sheet = book.active
    sheet.append(["系统SKU", "自定义SKU", "可用库存"])
    sheet.append(["1001", "10000001-red", 8])
    sheet.append(["2001", "child-special", 2])
    book.save(path)


def _presale_file(path):
    book = Workbook()
    sheet = book.active
    sheet.append(["说明"])
    sheet.append(["说明"])
    sheet.append(["商品ID", "规格编号", "出货天数"])
    sheet.append(["2", "700002", 7])
    book.save(path)


def test_worker_applies_combo_fallback_low_stock_zero_and_presale_bonus(tmp_path):
    sales = tmp_path / "sales.xlsx"
    stock = tmp_path / "stock.xlsx"
    presale = tmp_path / "presale.xlsx"
    combo = tmp_path / "combo.csv"
    _sales_file(sales)
    _stock_file(stock)
    _presale_file(presale)
    combo.write_bytes((
        "组合系统SKU\t组合自定义SKU\t子系统SKU\t子自定义SKU\tSkuNum\n"
        "9001\tzuhe-special\t2001\tchild-special\t1\n"
    ).encode("utf-16"))

    worker = AutoMatchWorker(
        str(sales), str(stock), str(presale),
        include_presale=True, combo_path=str(combo),
    )
    worker.run()

    assert len(worker.items) == 2
    low_stock = worker.items[0]
    assert low_stock["low_stock_zero"] is True
    assert low_stock["auto_status"] == "已确认"
    assert low_stock["auto_av"] == 0

    presale_item = worker.items[1]
    assert presale_item["combo_hit"] is True
    assert presale_item["presale_bonus"] == 9999
    assert presale_item["auto_av"] == 10001


def test_worker_can_disable_each_automatic_inventory_rule(tmp_path):
    sales = tmp_path / "sales.xlsx"
    stock = tmp_path / "stock.xlsx"
    presale = tmp_path / "presale.xlsx"
    _sales_file(sales)
    _stock_file(stock)
    _presale_file(presale)
    combo = tmp_path / "combo.csv"
    combo.write_bytes((
        "组合系统SKU\t组合自定义SKU\t子系统SKU\t子自定义SKU\tSkuNum\n"
        "9001\tzuhe-special\t2001\tchild-special\t1\n"
    ).encode("utf-16"))

    worker = AutoMatchWorker(
        str(sales), str(stock), str(presale),
        include_presale=True,
        presale_bonus_enabled=False,
        low_stock_zero_enabled=False,
        combo_path=str(combo),
    )
    worker.run()

    low_stock = worker.items[0]
    assert low_stock["low_stock_zero"] is False
    assert low_stock["auto_av"] == 8

    presale_item = worker.items[1]
    assert presale_item["presale_bonus"] == 0
    assert presale_item["auto_av"] == 2
