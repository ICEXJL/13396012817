# -*- coding: utf-8 -*-
"""销售资料异常 XML 的读取回归测试。"""

from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

import loader
from loader import load_sales_rows_robust
from loader import load_combo_sku_mapping_robust, resolve_combo_inventory


def _make_invalid_active_pane_xlsx(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品ID", "商品名称", "主商品货号", "商品货号", "卖家库存"])
    sheet.append(["1001", "测试商品", "10000001-black", "10000001-black", 3])
    workbook.save(path)

    fixed = path.with_name("invalid-pane.xlsx")
    with ZipFile(path, "r") as source, ZipFile(fixed, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                marker = b"<sheetView "
                insertion = b'<pane activePane="bottom" state="frozen" ySplit="1"/>'
                start = data.index(marker)
                end = data.index(b">", start) + 1
                data = data[:end] + insertion + data[end:]
            target.writestr(entry, data)
    return fixed


def test_loader_reads_workbook_with_invalid_active_pane(tmp_path):
    source = tmp_path / "source.xlsx"
    invalid = _make_invalid_active_pane_xlsx(source)

    rows = load_sales_rows_robust(str(invalid))

    assert len(rows) == 1
    assert rows[0]["cid"] == "1001"
    assert rows[0]["sheet"] == "Sheet"


def test_loader_reads_workbook_with_stale_dimension(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品ID", "商品名称", "主商品货号", "商品货号", "卖家库存"])
    sheet.append(["1002", "维度异常商品", "10000002-black", "10000002-black", 4])
    workbook.save(source)

    broken = tmp_path / "stale-dimension.xlsx"
    with ZipFile(source, "r") as source_zip, ZipFile(broken, "w", ZIP_DEFLATED) as target:
        for entry in source_zip.infolist():
            data = source_zip.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b'<dimension ref="A1:E2"></dimension>',
                    b'<dimension ref="A1"></dimension>',
                )
                data = data.replace(b'<dimension ref="A1:E2"/>', b'<dimension ref="A1"/>')
                data = data.replace(b'<dimension ref="A1:E2" />', b'<dimension ref="A1" />')
                assert (
                    b'<dimension ref="A1"></dimension>' in data
                    or b'<dimension ref="A1"/>' in data
                    or b'<dimension ref="A1" />' in data
                )
            target.writestr(entry, data)

    rows = load_sales_rows_robust(str(broken))

    assert len(rows) == 1
    assert rows[0]["cid"] == "1002"


def test_loader_reads_presale_specs_from_third_row_header(tmp_path):
    """出货天数表以第3行表头时，只返回出货天数大于1的规格编号。"""
    source = tmp_path / "出货天数.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["接口字段", "", "", "", "", "", "", "", "", ""])
    sheet.append(["接口数据", "", "", "", "", "", "", "", "", ""])
    sheet.append(["商品ID", "主商品货号", "商品名称", "规格编号", "名称", "类目", "非预购出货天数", "预购出货天数范围", "出货天数", "失败原因"])
    sheet.append(["1001", "A", "测试", "700001", "红", "类目", 1, "7-30", 1, ""])
    sheet.append(["1002", "A", "测试", "700002", "蓝", "类目", 1, "7-30", 7, ""])
    sheet.append(["1003", "A", "测试", "700003.0", "绿", "类目", 1, "7-30", "30", ""])
    sheet.append(["1004", "A", "测试", "", "黑", "类目", 1, "7-30", 30, ""])
    workbook.save(source)

    assert hasattr(loader, "load_presale_specs_robust")
    specs, meta = loader.load_presale_specs_robust(str(source))

    assert specs == {"700002": 7, "700003": 30}
    assert meta["sheet"] == "Sheet"
    assert meta["source_rows"] == 2


def test_sales_loader_exposes_specification_code(tmp_path):
    """销售资料的规格编号必须随行返回，用于识别预售单。"""
    source = tmp_path / "销售资料.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品ID", "商品名称", "规格编号", "主商品货号", "商品货号", "卖家库存"])
    sheet.append(["1001", "测试商品", "700002", "10000001-black", "10000001-black", 3])
    workbook.save(source)

    rows = load_sales_rows_robust(str(source))

    assert rows[0].get("spec") == "700002"


def test_tolerant_writer_updates_workbook_with_invalid_active_pane(tmp_path):
    """完成匹配时也应兼容异常 pane，且可写入库存列。"""
    source = tmp_path / "source.xlsx"
    invalid = _make_invalid_active_pane_xlsx(source)

    assert hasattr(loader, "write_sales_inventory_values")
    written = loader.write_sales_inventory_values(str(invalid), "Sheet", {2: 9})

    workbook = __import__("openpyxl").load_workbook(invalid, data_only=True)
    try:
        assert written == 1
        assert workbook["Sheet"].cell(row=2, column=9).value == 9
    finally:
        workbook.close()


def test_presale_loader_exposes_source_rows_and_writer_updates_days(tmp_path):
    source = tmp_path / "出货天数.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["说明", "", "", "", "", "", "", "", "", ""])
    sheet.append(["说明", "", "", "", "", "", "", "", "", ""])
    sheet.append(["商品ID", "主商品货号", "商品名称", "规格编号", "名称", "类目", "非预购出货天数", "预购出货天数范围", "出货天数", "失败原因"])
    sheet.append(["1001", "A", "测试", "700002", "红", "类目", 1, "7-30", 7, ""])
    workbook.save(source)

    specs, meta = loader.load_presale_specs_robust(str(source))
    assert specs == {"700002": 7}
    assert meta["rows_by_spec"] == {"700002": [4]}
    assert meta["days_col"] == 9

    assert loader.write_presale_days_values(
        str(source), meta["sheet"], meta["days_col"], {4: 1}
    ) == 1
    workbook = __import__("openpyxl").load_workbook(source, data_only=True)
    try:
        assert workbook[meta["sheet"]].cell(row=4, column=9).value == 1
    finally:
        workbook.close()


def test_combo_loader_reads_utf16_tsv_and_resolves_b_to_multiple_c_components(tmp_path):
    source = tmp_path / "组合SKU.csv"
    source.write_bytes((
        "组合系统SKU\t组合自定义SKU\t子系统SKU\t子自定义SKU\t子SKU产品名称\tSkuNum\n"
        "9001\tcombo-a\t1001\tchild-a\tA\t2\n"
        "9001\tcombo-a\t1002\tchild-b\tB\t1\n"
    ).encode("utf-16"))

    mapping, meta = load_combo_sku_mapping_robust(str(source))

    assert meta["rows"] == 2
    assert meta["unique_combos"] == 1
    assert [row["system_sku"] for row in mapping["combo-a"]] == ["1001", "1002"]
    result = resolve_combo_inventory("COMBO-A", mapping, {
        "1001": [("1001", "child-a", 9)],
        "1002": [("1002", "child-b", 8)],
    })
    assert result["status"] == "OK"
    assert result["available"] == 4  # min(floor(9 / 2), floor(8 / 1))
