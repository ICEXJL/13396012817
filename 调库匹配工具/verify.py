# -*- coding: utf-8 -*-
"""校验：核对写回文件与报告一致性 + 按匹配级别分布 + 抽查样本"""
import pandas as pd
from openpyxl import load_workbook
from collections import Counter
import re

XLSX = r"C:\Users\Administrator\Desktop\调库\Shopee-PH09-销售资料.xlsx"
CSV  = r"C:\Users\Administrator\Desktop\调库\调库匹配\调库匹配明细报告.csv"

rep = pd.read_csv(CSV, dtype=str).fillna("")

# 1) 读回写好的文件 I 列
wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["Sheet1"]
cur = {}
for r_idx, row in enumerate(ws.iter_rows(min_row=4, max_col=10), start=4):
    cid = row[0].value
    if cid is None or str(cid).strip() in ("", "nan"):
        continue
    cur[r_idx] = row[8].value
wb.close()

rep["Excel行号"] = rep["Excel行号"].astype(int)
# 报告行号与文件行号对应（报告是逐行处理，行号即Excel行号）
mismatch = 0
for _, r in rep.iterrows():
    rn = int(r["Excel行号"])
    if rn not in cur:
        print("!! 报告中行号不在文件里:", rn)
        mismatch += 1
        continue
    if r["匹配状态"] in ("已匹配", "已匹配(含无备货标注)", "部分匹配-需人工核对"):
        expect = r["写入卖家库存(I列)"]
        actual = cur[rn]
        if str(actual).strip() != str(expect).strip():
            print(f"!! 不一致 行{rn}: 报告={expect} 文件={actual} 货号={r['商品货号(F列)']}")
            mismatch += 1
print("写回校验: 不一致数 =", mismatch)

# 2) 按匹配级别分布（从说明里提取 L0/L1/L2... 或 E列兜底）
def level_of(desc, st):
    if "未匹配" == st:
        return "未匹配"
    if "多候选" in st:
        return "多候选"
    if "部分匹配" in st:
        return "部分匹配"
    if "E列兜底" in desc:
        return "E列兜底"
    m = re.search(r"(L[0-8])", desc)
    return m.group(1) if m else "其他"

rep["级别"] = rep.apply(lambda r: level_of(r["匹配规则说明"], r["匹配状态"]), axis=1)
print("\n== 按匹配级别统计(行数) ==")
for k, v in Counter(rep["级别"]).most_common():
    print(f"  {k}: {v}")

# 3) 抽查各级样本（各3条）
print("\n== 抽查样本 ==")
for lvl in ["L1", "L2", "L3", "L4", "L5", "L6", "L7", "E列兜底"]:
    sub = rep[rep["级别"] == lvl]
    if len(sub) == 0:
        continue
    print(f"\n-- {lvl} (共{len(sub)}行) 样本:")
    for _, r in sub.head(3).iterrows():
        print(f"   {r['商品货号(F列)']!r} -> {r['匹配到的库存SKU']!r} 可用={r['可用库存']} 写入={r['写入卖家库存(I列)']}")
        print(f"       {r['匹配规则说明'][:120]}")

# 4) 未匹配/多候选/部分匹配 详细分类
print("\n== 待人工处理分类 ==")
pending = rep[rep["匹配状态"].isin(["未匹配", "多候选-未写入", "部分匹配-需人工核对"])]
cat = Counter()
for _, r in pending.iterrows():
    f = r["商品货号(F列)"]
    if not re.search(r"\d", f):
        cat["F列无数字(颜色/描述等脏数据)"] += 1
    elif "多候选" in r["匹配状态"]:
        cat["多候选(规范化后歧义)"] += 1
    elif "部分匹配" in r["匹配状态"]:
        cat["+号组合部分命中"] += 1
    elif re.search(r"[（(]\d+", f):
        cat["括号内尺寸/数量无法对应"] += 1
    elif re.search(r"[x×*]|cm|mm|ml|m\b", f, re.I):
        cat["含尺寸规格无法对应"] += 1
    elif "无备货" in f:
        cat["无备货标注但基础SKU未找到"] += 1
    else:
        cat["其他"] += 1
for k, v in cat.most_common():
    print(f"  {k}: {v}")
