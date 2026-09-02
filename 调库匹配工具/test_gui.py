# -*- coding: utf-8 -*-
"""test_gui.py —— 无头 GUI 集成测试：构造窗口、模拟匹配/决策/写回"""
import os
import sys
import shutil
import tempfile
from pathlib import Path

os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.path.insert(0, r"C:\Users\Administrator\Desktop\调库\调库匹配")

from PySide6.QtWidgets import QApplication, QMessageBox, QInputDialog
from match_stock_gui import MainWindow, AutoMatchWorker
from openpyxl import load_workbook

WORK_DIR = Path(r"C:\Users\Administrator\Desktop\调库")
SALES = str(WORK_DIR / "Shopee-PH09-销售资料.xlsx")
STOCK = str(max(WORK_DIR.glob("导出当前条件所有库存记录*.xls"),
                key=lambda path: path.stat().st_mtime))
PRESALE = str(WORK_DIR / "Shopee109-出货天数.xlsx")

# 用副本测试写回，避免动原文件
tmp = tempfile.mkdtemp(prefix="gui_test_")
copy = os.path.join(tmp, "销售测试.xlsx")
shutil.copy2(SALES, copy)

app = QApplication([])
w = MainWindow()
print("[1] 窗口构造成功")
assert w.btn_adopt.text() == "采用自动匹配值"
assert w.btn_keep.text() == "保留原值"
assert w.btn_zero.text() == "调为 0"
assert w.btn_custom.text() == "手动输入库存"
assert w.listw is not None and w.cards_grid is not None

# 直接设置路径（跳过文件对话框）
w.sales_path = copy
w.stock_path = STOCK
w.presale_path = PRESALE
w.lbl_sales.setText("测试")
w.lbl_stock.setText("测试")
w.lbl_presale.setText("测试")

# 同步执行 worker 逻辑
worker = AutoMatchWorker(copy, STOCK, PRESALE)
worker.run()
items, meta = worker.items, f"ds | {Path(STOCK).name}"
print(f"[2] worker 完成: {len(items)} 行, meta={meta}")
assert len(items) > 0, f"销售资料没有可处理行: {len(items)}"

w._on_match_done(items, meta, "Sheet1 | 测试")
print("[3] 进入复核模式, 当前idx:", w.idx, "| 待处理:", sum(1 for i in items if i['decision'] is None and i['auto_status'] in ('建议','待手选')))

# ---- 模拟决策（注意：每次决策后 idx 会自动前进到下一个未处理） ----
def goto(status=None, undecided=False):
    """跳到满足条件的行"""
    for _ in range(7000):
        it = w.items[w.idx]
        if status and it["auto_status"] != status:
            w.idx = (w.idx + 1) % len(w.items); continue
        if undecided and it["decision"]:
            w.idx = (w.idx + 1) % len(w.items); continue
        return it
    return None

# 待手选行: 匹配为0
it0 = goto("待手选", True)
f0 = it0["f"]
w._decide_zero()
assert it0["decision"] == "调为0" and it0["user_av"] == 0
print(f"[4] 匹配为0: {f0!r} -> OK")

# 待手选行: 选卡片确认
it1 = goto("待手选", True)
if it1 and it1["candidates"]:
    card = next((c for c in w.card_widgets if hasattr(c, "sku")), None)
    assert card is not None, "应有候选卡片"
    w._on_card_clicked(card)
    w._decide_adopt()
    assert it1["decision"] == "采用自动值" and it1["user_av"] == card.av
    print(f"[5] 确认卡片: {it1['f']!r} -> {card.sku} 库存={it1['user_av']}")

# 自定义数值
it2 = goto(undecided=True)
QInputDialog.getInt = lambda *a, **k: (321, True)
w._decide_custom()
assert it2["decision"] == "手动输入" and it2["user_av"] == 321
print(f"[6] 自定义: {it2['f']!r} -> {it2['user_av']}")

# 保留原值
it3 = goto(undecided=True)
w._decide_keep()
assert it3["decision"] == "保留原值"
print(f"[7] 跳过: {it3['f']!r} -> OK")

# ---- 写回 ----
import match_stock_gui
match_stock_gui.clear_state = lambda: None   # 沙箱禁止删除文件，mock掉（真实环境无此限制）
QMessageBox.question = lambda *a, **k: QMessageBox.Yes
QMessageBox.exec = lambda self: 0
w.finish_writeback()
print("[8] 写回完成")

# 验证写回结果
wb = load_workbook(copy, data_only=True)
ws = wb.active
checked = 0
for it in items:
    d = it["decision"]
    if d == "调为0":
        assert ws.cell(row=it["row"], column=9).value == 0, f"行{it['row']}应为0"
        checked += 1
    elif d == "手动输入":
        assert ws.cell(row=it["row"], column=9).value == 321, f"行{it['row']}应为321"
        checked += 1
    elif d == "采用自动值":
        assert ws.cell(row=it["row"], column=9).value == it["user_av"], f"行{it['row']}"
        checked += 1
    elif it["auto_status"] == "已确认" and d is None:
        assert ws.cell(row=it["row"], column=9).value == it["auto_av"], f"行{it['row']}自动确认"
        checked += 1
print(f"[9] 写回校验通过: 检查了 {checked} 个单元格")

# 报告字段与状态文件清理
import pandas as pd
report_df = pd.read_csv(w._last_report, encoding="utf-8-sig")
assert {"规格编号", "是否预售单", "出货天数", "复核决定", "最终写入库存", "原卖家库存"}.issubset(report_df.columns)
keep_row = next(it for it in items if it["decision"] == "保留原值")
assert ws.cell(row=keep_row["row"], column=9).value == keep_row["i_old"]
from review_state import _state_path
print("[10] 状态文件存在:", os.path.exists(_state_path()))
w.close()
print("全部 GUI 集成测试通过")
